"""
Pattern for assembling the final .xlsx deliverable from reconciled row data plus the reference
sheets pulled from the original (usually .xlsb) target workbook.

Typical flow this supports:
  1. Read the original .xlsb with pyxlsb (read-only) to pull out reference sheets verbatim --
     Dictionary, Validations, geoid-crosswalk, or whatever equivalents exist -- as plain
     list-of-lists / list-of-dicts data (pyxlsb can't preserve formulas/formatting, only values).
  2. Build the populated data sheet (e.g. "Capabilities") from your reconciled row dicts with
     openpyxl, using the target schema's exact column order.
  3. Write everything into one new .xlsx workbook and hand that back as the deliverable, since
     .xlsb cannot be written directly by these tools.

Adapt TARGET_COLS and the reference-sheet list to your actual schema -- this is a structural
pattern, not a fixed schema.
"""

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter


def build_workbook(target_cols, final_rows, reference_sheets, out_path,
                    data_sheet_name='Capabilities', column_widths=None):
    """
    target_cols: ordered list of column headers for the main data sheet.
    final_rows: list of dicts, each keyed by (a superset of) target_cols. Missing/blank values
                are written as empty cells rather than the string 'None'.
    reference_sheets: dict of {sheet_name: list_of_rows}, where list_of_rows is a list of lists
                      (first row treated as header and bolded). Pass sheets pulled verbatim from
                      the original workbook (Dictionary, Validations, geoid-crosswalk, etc.).
    out_path: where to save the resulting .xlsx.
    column_widths: optional {column_header: width} overrides for the data sheet; anything not
                   listed gets a sane default.
    """
    wb = Workbook()

    ws = wb.active
    ws.title = data_sheet_name
    header_font = Font(name='Arial', bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    body_font = Font(name='Arial')

    for j, col in enumerate(target_cols, start=1):
        cell = ws.cell(row=1, column=j, value=col.strip())
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    ws.freeze_panes = 'A2'
    ws.row_dimensions[1].height = 30

    for i, row in enumerate(final_rows, start=2):
        for j, col in enumerate(target_cols, start=1):
            val = row.get(col, '')
            if val == '':
                val = None
            c = ws.cell(row=i, column=j, value=val)
            c.font = body_font

    widths = column_widths or {}
    for j, col in enumerate(target_cols, start=1):
        ws.column_dimensions[get_column_letter(j)].width = widths.get(col.strip(), 16)

    for sheet_name, rows in reference_sheets.items():
        rs = wb.create_sheet(sheet_name)
        for i, row in enumerate(rows, start=1):
            for j, val in enumerate(row, start=1):
                c = rs.cell(row=i, column=j, value=val)
                c.font = Font(name='Arial', bold=(i == 1))

    wb.save(out_path)
    return out_path


# --- Example: pulling reference sheets out of the original .xlsb with pyxlsb -------------------
def extract_xlsb_sheet_as_rows(xlsb_path, sheet_name):
    """Read one sheet of a .xlsb file into a plain list-of-lists (values only, no formulas)."""
    import pyxlsb
    with pyxlsb.open_workbook(xlsb_path) as wb:
        with wb.get_sheet(sheet_name) as sheet:
            return [[c.v for c in r] for r in sheet.rows()]


def extract_crosswalk_as_dicts(xlsb_path, sheet_name='geoid-crosswalk'):
    """Read a crosswalk-style sheet into a list of dicts keyed by its header row."""
    rows = extract_xlsb_sheet_as_rows(xlsb_path, sheet_name)
    header = rows[0]
    return [dict(zip(header, r)) for r in rows[1:] if r and r[0] is not None]
