"""
Build the NYSDOT PPDAF Final Report .docx from the v1 HTML.

Not a generic HTML->Word converter. It walks the report's own known structure
(TransportNY design-system v2 document shell) and emits Word constructs:
title page, disclaimer, the DOT F 1700.7 grid as a real Word table, a TOC
field, Heading 1/2, body text, bullets, tables with captions, and figures with
numbered captions. Design chrome (kickers, rules, card grids used purely for
layout) is dropped; card grids that carry content become tables or headed
blocks.

Run:  python scratchpad/nysdot-final-report/html2docx.py
Out:  src/themes/.../reports/Final_report_draft_v1.docx
"""
import os, re, sys
from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..'))
REPORTS = os.path.join(ROOT, 'src', 'themes', 'transportny',
                       'TransportNY Design System', 'dms_design_system_v2', 'reports')
SRC = os.path.join(REPORTS, 'nysdot-ppdaf-final-report-v1.html')
OUT = os.path.join(REPORTS, 'Final_report_draft_v1.docx')

INK    = RGBColor(0x0F, 0x2D, 0x4D)   # TransportNY deep navy
SLATE  = RGBColor(0x33, 0x41, 0x55)
MUTED  = RGBColor(0x64, 0x74, 0x8B)
ACCENT = RGBColor(0xCA, 0x8A, 0x04)   # TransportNY gold
RULE   = 'D9DEE5'

fig_n = 0   # figures are numbered in the HTML captions already; we reuse those

SECTION_NUMBERS = {'s1': '1', 's2': '2', 's3': '3', 's4': '4',
                   's5': '5', 's6': '6', 's7': '7', 's8': '8'}


# ── low-level helpers ────────────────────────────────────────────────────────
def shade(cell, hexcolor):
    el = OxmlElement('w:shd'); el.set(qn('w:val'), 'clear')
    el.set(qn('w:color'), 'auto'); el.set(qn('w:fill'), hexcolor)
    cell._tc.get_or_add_tcPr().append(el)


def borders(table, color=RULE, sz=4):
    tblPr = table._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement('w:' + edge)
        e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(sz))
        e.set(qn('w:space'), '0'); e.set(qn('w:color'), color)
        b.append(e)
    tblPr.append(b)


def keep_with_next(p):
    p.paragraph_format.keep_with_next = True


def field(paragraph, instr):
    """Insert a Word field code (used for the TOC and page numbers)."""
    r = paragraph.add_run()
    fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), 'begin')
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = instr
    sep = OxmlElement('w:fldChar'); sep.set(qn('w:fldCharType'), 'separate')
    t = OxmlElement('w:t'); t.text = 'Right-click and select Update Field.'
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
    r._r.append(fc); r._r.append(it); r._r.append(sep); r._r.append(t); r._r.append(end)


def hrule(doc):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr'); bot = OxmlElement('w:bottom')
    bot.set(qn('w:val'), 'single'); bot.set(qn('w:sz'), '6')
    bot.set(qn('w:space'), '4'); bot.set(qn('w:color'), RULE)
    pbdr.append(bot); pPr.append(pbdr)
    p.paragraph_format.space_after = Pt(10)
    return p


# ── inline text with bold / italic / code preserved ──────────────────────────
def add_inline(par, node, base_bold=False, base_italic=False, size=None):
    if isinstance(node, NavigableString):
        txt = re.sub(r'\s+', ' ', str(node))
        if not txt:
            return
        r = par.add_run(txt)
        r.bold = base_bold or None
        r.italic = base_italic or None
        if size: r.font.size = size
        return
    if not isinstance(node, Tag):
        return
    if node.name == 'br':
        par.add_run().add_break()
        return
    cls = node.get('class') or []
    bold = base_bold or node.name in ('b', 'strong')
    ital = base_italic or node.name in ('i', 'em')
    mono = 'tny-mono' in cls
    if mono:
        r = par.add_run(re.sub(r'\s+', ' ', node.get_text()))
        r.font.name = 'Consolas'; r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x37, 0x57, 0x6B)
        return
    for child in node.children:
        add_inline(par, child, bold, ital, size)


def text_of(node):
    return re.sub(r'\s+', ' ', node.get_text(' ', strip=True)).strip()


# ── block emitters ───────────────────────────────────────────────────────────
def tidy(p):
    """Strip the leading/trailing space that HTML indentation introduces."""
    if p.runs:
        p.runs[0].text = p.runs[0].text.lstrip()
        p.runs[-1].text = p.runs[-1].text.rstrip()
    return p


def body_par(doc, node, style='TNYBody'):
    p = doc.add_paragraph(style=style)
    for c in node.children:
        add_inline(p, c)
    return tidy(p)


def bullet(doc, node):
    p = doc.add_paragraph(style='TNYBullet')
    for c in node.children:
        add_inline(p, c)
    return tidy(p)


def emit_table(doc, tbl, caption=None):
    rows = tbl.find_all('tr')
    if not rows:
        return
    ncol = 0
    for r in rows:
        n = sum(int(c.get('colspan', 1)) for c in r.find_all(['td', 'th']))
        ncol = max(ncol, n)
    if caption:
        cp = doc.add_paragraph(caption, style='TNYTableCaption')
        keep_with_next(cp)
    t = doc.add_table(rows=0, cols=ncol)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(t)
    for r in rows:
        cells_src = r.find_all(['td', 'th'])
        if not cells_src:
            continue
        row = t.add_row()
        is_head = cells_src[0].name == 'th'
        is_total = 'total' in (r.get('class') or [])
        ci = 0
        for c in cells_src:
            span = int(c.get('colspan', 1))
            if ci >= ncol:
                break
            cell = row.cells[ci]
            if span > 1 and ci + span <= ncol:
                cell = cell.merge(row.cells[ci + span - 1])
            cell.text = ''
            p = cell.paragraphs[0]
            p.style = doc.styles['TNYCell']
            for ch in c.children:
                add_inline(p, ch, base_bold=is_head or is_total)
            if is_head:
                shade(cell, 'F1F4F7')
                for run in p.runs:
                    run.font.size = Pt(8); run.font.color.rgb = MUTED
                    run.bold = True
            elif is_total:
                shade(cell, 'F7F8FA')
            ci += span
    doc.add_paragraph(style='TNYSpacer')
    return t


def emit_figure(doc, figure):
    img = figure.find('img')
    cap = figure.find('figcaption')
    if img is not None:
        path = os.path.normpath(os.path.join(REPORTS, img['src']))
        if os.path.exists(path):
            with Image.open(path) as im:
                w, h = im.size
            width = Inches(6.3)
            p = doc.add_paragraph(style='TNYFigure')
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(path, width=width)
            keep_with_next(p)
        else:
            doc.add_paragraph('[missing image: %s]' % img['src'], style='TNYBody')
    if cap is not None:
        p = doc.add_paragraph(style='TNYFigCaption')
        for c in cap.children:
            add_inline(p, c)
    doc.add_paragraph(style='TNYSpacer')


def emit_stat_cards(doc, container, heading=None):
    """A grid of .tny-card blocks becomes a two-column Word table so the
    numbers keep their prominence without depending on CSS grid."""
    cards = [d for d in container.find_all('div', recursive=False)
             if 'tny-card' in (d.get('class') or [])]
    if not cards:
        return False
    t = doc.add_table(rows=0, cols=2)
    t.style = 'Table Grid'
    borders(t, color='E6EAEF')
    for i in range(0, len(cards), 2):
        row = t.add_row()
        for j in range(2):
            cell = row.cells[j]
            cell.text = ''
            if i + j >= len(cards):
                continue
            card = cards[i + j]
            first = True
            for child in card.children:
                if not isinstance(child, Tag):
                    continue
                txt = text_of(child)
                if not txt:
                    continue
                cls = ' '.join(child.get('class') or [])
                p = cell.paragraphs[0] if first else cell.add_paragraph()
                first = False
                p.style = doc.styles['TNYCell']
                if 'text-[34px]' in cls or 'text-[27px]' in cls:
                    r = p.add_run(txt); r.bold = True
                    r.font.size = Pt(20); r.font.color.rgb = INK
                elif 'font-display' in cls:
                    r = p.add_run(txt); r.bold = True; r.font.size = Pt(10.5)
                    r.font.color.rgb = INK
                elif 'chip' in cls:
                    r = p.add_run(txt.upper()); r.font.size = Pt(7.5)
                    r.font.color.rgb = MUTED
                else:
                    for ch in child.children:
                        add_inline(p, ch)
                    for r in p.runs:
                        r.font.size = Pt(8.5)
    doc.add_paragraph(style='TNYSpacer')
    return True


# ── document styles ──────────────────────────────────────────────────────────
def build_styles(doc):
    st = doc.styles
    normal = st['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = SLATE
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.15

    def mk(name, base='Normal'):
        try:
            return st[name]
        except KeyError:
            return st.add_style(name, 1)  # WD_STYLE_TYPE.PARAGRAPH

    s = mk('TNYBody'); s.base_style = st['Normal']
    s.font.size = Pt(10.5); s.paragraph_format.space_after = Pt(8)
    s.paragraph_format.line_spacing = 1.25

    s = mk('TNYBullet'); s.base_style = st['List Bullet']
    s.font.size = Pt(10.5); s.font.color.rgb = SLATE
    s.paragraph_format.space_after = Pt(4)
    s.paragraph_format.line_spacing = 1.2

    s = mk('TNYLead'); s.base_style = st['Normal']
    s.font.size = Pt(11.5); s.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
    s.paragraph_format.space_after = Pt(12)
    s.paragraph_format.line_spacing = 1.3

    s = mk('TNYCell'); s.base_style = st['Normal']
    s.font.size = Pt(9); s.paragraph_format.space_after = Pt(2)
    s.paragraph_format.line_spacing = 1.1

    s = mk('TNYTableCaption'); s.base_style = st['Normal']
    s.font.size = Pt(8.5); s.font.bold = True; s.font.color.rgb = MUTED
    s.font.all_caps = True
    s.paragraph_format.space_before = Pt(10); s.paragraph_format.space_after = Pt(4)

    s = mk('TNYFigCaption'); s.base_style = st['Normal']
    s.font.size = Pt(8.5); s.font.color.rgb = MUTED
    s.paragraph_format.space_after = Pt(6)
    s.paragraph_format.line_spacing = 1.15

    s = mk('TNYFigure'); s.base_style = st['Normal']
    s.paragraph_format.space_before = Pt(8); s.paragraph_format.space_after = Pt(3)

    s = mk('TNYSpacer'); s.base_style = st['Normal']
    s.font.size = Pt(4); s.paragraph_format.space_after = Pt(0)

    s = mk('TNYCallout'); s.base_style = st['Normal']
    s.font.size = Pt(10); s.paragraph_format.space_after = Pt(6)
    s.paragraph_format.left_indent = Inches(0.25)
    s.paragraph_format.line_spacing = 1.25

    h1 = st['Heading 1']
    h1.font.name = 'Calibri Light'; h1.font.size = Pt(19); h1.font.bold = True
    h1.font.color.rgb = INK
    h1.paragraph_format.space_before = Pt(20); h1.paragraph_format.space_after = Pt(8)
    h1.paragraph_format.keep_with_next = True

    h2 = st['Heading 2']
    h2.font.name = 'Calibri Light'; h2.font.size = Pt(13); h2.font.bold = True
    h2.font.color.rgb = INK
    h2.paragraph_format.space_before = Pt(14); h2.paragraph_format.space_after = Pt(5)
    h2.paragraph_format.keep_with_next = True

    h3 = st['Heading 3']
    h3.font.name = 'Calibri'; h3.font.size = Pt(11); h3.font.bold = True
    h3.font.color.rgb = RGBColor(0x37, 0x57, 0x6B)
    h3.paragraph_format.space_before = Pt(10); h3.paragraph_format.space_after = Pt(4)
    h3.paragraph_format.keep_with_next = True


def page_setup(doc):
    for sec in doc.sections:
        sec.page_width = Inches(8.5); sec.page_height = Inches(11)
        sec.left_margin = sec.right_margin = Inches(1.0)
        sec.top_margin = Inches(0.9); sec.bottom_margin = Inches(0.9)
        # footer: title left, page number right
        f = sec.footer.paragraphs[0]
        f.text = ''
        r = f.add_run('PPDAF Final Report  ·  Task Assignment SP-20-03  ·  NYSDOT × AVAIL\t\t')
        r.font.size = Pt(8); r.font.color.rgb = MUTED
        field(f, 'PAGE')
        for rr in f.runs:
            if rr.font.size is None:
                rr.font.size = Pt(8); rr.font.color.rgb = MUTED


# ── section walkers ──────────────────────────────────────────────────────────
def walk_generic(doc, root):
    """Emit the content of a report <section> in document order."""
    for el in root.children:
        if not isinstance(el, Tag):
            continue
        emit_node(doc, el)


def emit_node(doc, el):
    cls = el.get('class') or []
    clsx = ' '.join(cls)

    if el.name == 'h2':
        doc.add_heading(text_of(el), level=1); return
    if el.name == 'h3':
        doc.add_heading(text_of(el), level=2); return
    if el.name == 'h4':
        doc.add_heading(text_of(el), level=3); return
    if el.name == 'figure':
        emit_figure(doc, el); return
    if el.name == 'table':
        cap = el.find('caption')
        emit_table(doc, el, text_of(cap) if cap else None); return

    # section lede paragraph (direct <p> under a section)
    if el.name == 'p':
        if 'kicker' in clsx or not text_of(el):
            return
        body_par(doc, el, 'TNYLead' if 'text-[14px]' in clsx else 'TNYBody')
        return

    if el.name in ('ul', 'ol'):
        for li in el.find_all('li', recursive=False):
            bullet(doc, li)
        doc.add_paragraph(style='TNYSpacer')
        return

    # the kicker rule strip: drop
    if 'flex' in clsx and el.find('span', class_='kicker'):
        return

    # prose blocks
    if 'prose-tny' in clsx:
        for child in el.children:
            if isinstance(child, Tag):
                emit_node(doc, child)
        return

    # a card carrying a caption'd table
    if 'tny-card' in clsx:
        tbl = el.find('table')
        ink = 'tny-card-ink' in clsx
        bone = 'tny-card-bone' in clsx
        if tbl is not None and not ink:
            chip = el.find('div', class_=re.compile('chip'))
            cap_el = tbl.find('caption')
            cap = text_of(cap_el) if cap_el else (text_of(chip) if chip else None)
            emit_table(doc, tbl, cap)
            for sib in el.find_all('p', recursive=False):
                if text_of(sib):
                    body_par(doc, sib, 'TNYFigCaption')
            return
        # callout card (ink / bone): heading chip + prose
        chip = el.find('div', class_=re.compile('chip'))
        if chip is not None and text_of(chip):
            p = doc.add_paragraph(style='TNYTableCaption')
            r = p.add_run(text_of(chip)); r.font.color.rgb = ACCENT
            keep_with_next(p)
        for child in el.children:
            if not isinstance(child, Tag):
                continue
            if child is chip:
                continue
            if child.name == 'p':
                body_par(doc, child, 'TNYCallout')
            elif child.name in ('ul', 'ol'):
                for li in child.find_all('li', recursive=False):
                    bullet(doc, li)
            elif child.name == 'table':
                emit_table(doc, child, None)
        doc.add_paragraph(style='TNYSpacer')
        return

    # grids of cards → two-column table; grids of other things → recurse
    if 'grid' in clsx:
        if emit_stat_cards(doc, el):
            return
        for child in el.children:
            if isinstance(child, Tag):
                emit_node(doc, child)
        return

    # anything else: recurse
    for child in el.children:
        if isinstance(child, Tag):
            emit_node(doc, child)


# ── front matter, hand-built ─────────────────────────────────────────────────
def title_page(doc, soup):
    hdr = soup.find('header')
    h1 = hdr.find('h1')

    for _ in range(3):
        doc.add_paragraph(style='TNYSpacer')

    p = doc.add_paragraph()
    r = p.add_run('SR-21-07  ·  TASK ASSIGNMENT SP-20-03')
    r.font.size = Pt(9); r.font.bold = True; r.font.color.rgb = ACCENT
    p.paragraph_format.space_after = Pt(2)

    p = doc.add_paragraph()
    r = p.add_run('FINAL REPORT')
    r.font.size = Pt(9); r.font.bold = True; r.font.color.rgb = MUTED
    p.paragraph_format.space_after = Pt(18)

    p = doc.add_paragraph()
    r = p.add_run('Research, Development and Support of an\n'
                  'Integrated Planning and Performance Data\n'
                  'and Analytics Framework (PPDAF)')
    r.font.size = Pt(26); r.font.bold = True; r.font.name = 'Calibri Light'
    r.font.color.rgb = INK
    p.paragraph_format.space_after = Pt(14)
    p.paragraph_format.line_spacing = 1.1

    lead = hdr.find('p')
    p = doc.add_paragraph(style='TNYLead')
    for c in lead.children:
        add_inline(p, c)

    # cover image
    cover = os.path.normpath(os.path.join(REPORTS, '../assets/screens/macro-01-overview.png'))
    if os.path.exists(cover):
        cp = doc.add_paragraph(style='TNYFigure')
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.add_run().add_picture(cover, width=Inches(6.3))
        cap = doc.add_paragraph(style='TNYFigCaption')
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.add_run('TransportNY NPMRDS Macro View — statewide Level of Travel Time Reliability, CY 2025.')

    # prepared for / by / through
    t = doc.add_table(rows=1, cols=3)
    borders(t, color='FFFFFF', sz=0)
    cells = t.rows[0].cells
    blocks = [
        ('PREPARED FOR',
         'New York State Department of Transportation\nOffice of Policy, Planning & Performance\n'
         '50 Wolf Road, Albany, NY 12232'),
        ('PREPARED BY',
         'Catherine T. Lawson, Ph.D., Principal Investigator\n'
         'Albany Visualization and Informatics Lab (AVAIL)\n'
         'Lewis Mumford Center, University at Albany, SUNY\n'
         '1400 Washington Avenue, Albany, NY 12222'),
        ('THROUGH',
         'The Research Foundation for SUNY, subconsultant to\n'
         'Rutgers, The State University of New Jersey\n'
         'Center for Advanced Infrastructure & Transportation\n'
         'Region 2 UTC Consortium'),
    ]
    for cell, (label, body) in zip(cells, blocks):
        cell.text = ''
        p = cell.paragraphs[0]; p.style = doc.styles['TNYCell']
        r = p.add_run(label); r.bold = True; r.font.size = Pt(7.5); r.font.color.rgb = MUTED
        p2 = cell.add_paragraph(style='TNYCell')
        r = p2.add_run(body); r.font.size = Pt(8.5)

    p = doc.add_paragraph()
    r = p.add_run('\nDraft v1  ·  September 2026  ·  NYSDOT Contract C000799  ·  '
                  'Federal Aid Project SP20(038)  ·  CFDA 20.205\n'
                  'Period of performance: 26 June 2021 – 30 September 2026')
    r.font.size = Pt(8.5); r.font.color.rgb = MUTED
    doc.add_page_break()


def disclaimer_page(doc, soup):
    doc.add_heading('Disclaimer', level=1)
    txt = ('This report was funded in part through grant(s) from the Federal Highway Administration, '
           'United States Department of Transportation, under the State Planning and Research Program, '
           'Section 505 of Title 23, U.S. Code. The contents of this report do not necessarily reflect '
           'the official views or policy of the United States Department of Transportation, the Federal '
           'Highway Administration or the New York State Department of Transportation. This report does '
           'not constitute a standard, specification, regulation, product endorsement, or an endorsement '
           'of manufacturers.')
    doc.add_paragraph(txt, style='TNYBody')
    doc.add_page_break()


def toc_page(doc):
    doc.add_heading('Table of Contents', level=1)
    p = doc.add_paragraph()
    field(p, r'TOC \o "1-2" \h \z \u')
    doc.add_paragraph()
    n = doc.add_paragraph(style='TNYFigCaption')
    n.add_run('This table of contents is a Word field. To populate it, select it, right-click, '
              'and choose Update Field → Update entire table.')
    doc.add_page_break()


# ── main ─────────────────────────────────────────────────────────────────────
def main():
    soup = BeautifulSoup(open(SRC, encoding='utf-8').read(), 'lxml')
    doc = Document()
    build_styles(doc)
    page_setup(doc)

    title_page(doc, soup)
    disclaimer_page(doc, soup)

    wrapper = soup.find('div', class_=re.compile(r'max-w-\[1180px\]'))
    sections = wrapper.find_all('section', recursive=False)

    # section 0 = disclaimer (handled), 1 = DOT F 1700.7, 2 = about+contents
    for sec in sections:
        sid = sec.get('id')
        h2 = sec.find('h2')
        title = text_of(h2) if h2 else ''

        if title == 'Disclaimer' or (h2 is None and sec.find('div', class_='tny-card-bone')):
            continue
        if title == 'Technical Report Documentation Page':
            doc.add_heading(title, level=1)
            doc.add_paragraph(
                'Form DOT F 1700.7 (8-72). Reproduction of completed page authorized.',
                style='TNYFigCaption')
            emit_table(doc, sec.find('table'), None)
            doc.add_page_break()
            continue
        if title == 'About this Document':
            doc.add_heading('About this Document', level=1)
            for pr in sec.find_all('div', class_='prose-tny'):
                for child in pr.children:
                    if isinstance(child, Tag):
                        emit_node(doc, child)
            toc_page(doc)
            continue

        # every remaining section starts on a new page
        if sid in ('exec', 's1', 's2', 's3', 's4', 's5', 's6', 's7', 's8', 'appx'):
            if sid == 'exec':
                doc.add_heading('Executive Summary', level=1)
                sub = doc.add_paragraph()
                r = sub.add_run(title)
                r.font.size = Pt(13); r.bold = True
                r.font.color.rgb = RGBColor(0x37, 0x57, 0x6B)
                sub.paragraph_format.space_after = Pt(10)
                h2.decompose()
            elif sid in SECTION_NUMBERS and h2 is not None:
                h2.string = '%s.  %s' % (SECTION_NUMBERS[sid], title)
            walk_generic(doc, sec)
            if sid != 'appx':
                doc.add_page_break()
        else:
            walk_generic(doc, sec)

    doc.save(OUT)
    size = os.path.getsize(OUT) / 1024 / 1024
    print('wrote %s  (%.2f MB)' % (OUT, size))
    print('paragraphs: %d   tables: %d   inline shapes: %d'
          % (len(doc.paragraphs), len(doc.tables), len(doc.inline_shapes)))


if __name__ == '__main__':
    main()
